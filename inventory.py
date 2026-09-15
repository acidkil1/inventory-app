import customtkinter as ctk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime, date
import sqlite3
from db_manager import get_connection, init_db
import openpyxl
import os
ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")
class InventoryApp(ctk.CTk):
    def __init__(self):
        super().__init__()
        icon_path = os.path.join(os.path.dirname(__file__), 'inventory.ico')
        if os.path.exists(icon_path):
            self.iconbitmap(icon_path)
        self.title("Система учета инвентаризации НФ УУНиТ")
        self.geometry("1200x700")
        init_db()
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)
        self.create_navigation()
        self.main_container = ctk.CTkFrame(self)
        self.main_container.grid(row=0, column=1, sticky="nsew", padx=10, pady=10)
        self.main_container.grid_columnconfigure(0, weight=1)
        self.main_container.grid_rowconfigure(0, weight=1)
        self.frames = {}
        self.create_frames()
        self.show_frame("dashboard")
    def create_navigation(self):
        nav_frame = ctk.CTkFrame(self, width=200, corner_radius=0)
        nav_frame.grid(row=0, column=0, sticky="nsew")
        nav_frame.grid_rowconfigure(8, weight=1)
        title_label = ctk.CTkLabel(
            nav_frame,
            text="ИНВЕНТАРИЗАЦИЯ",
            font=ctk.CTkFont(size=14, weight="bold")
        )
        title_label.grid(row=0, column=0, pady=20, padx=10)
        buttons = [
            ("Главная", "dashboard"),
            ("Добавить оборудование", "add"),
            ("Перемещение", "move"),
            ("Списание", "writeoff"),
            ("История", "history"),
            ("Отчеты", "reports"),
            ("Экспорт/Импорт", "import_export"),
        ]
        for i, (text, frame_name) in enumerate(buttons, start=1):
            btn = ctk.CTkButton(
                nav_frame,
                text=text,
                command=lambda f=frame_name: self.show_frame(f),
                width=180,
                height=40
            )
            btn.grid(row=i, column=0, pady=5, padx=10)
        exit_btn = ctk.CTkButton(
            nav_frame,
            text="Выход",
            command=self.quit,
            fg_color="red",
            hover_color="darkred",
            width=180,
            height=40
        )
        exit_btn.grid(row=9, column=0, pady=20, padx=10)
    def create_frames(self):
        frames_to_create = [
            ("dashboard", DashboardFrame),
            ("add", AddEquipmentFrame),
            ("move", MoveEquipmentFrame),
            ("writeoff", WriteOffFrame),
            ("history", HistoryFrame),
            ("reports", ReportsFrame),
            ("import_export", ImportExportFrame)
        ]
        for name, FrameClass in frames_to_create:
            frame = FrameClass(self.main_container, self)
            self.frames[name] = frame
            frame.grid(row=0, column=0, sticky="nsew")
    def show_frame(self, name):
        frame = self.frames.get(name)
        if frame:
            frame.tkraise()
            if name == "dashboard":
                frame.update_dashboard()
            elif name == "history":
                frame.load_history()
            elif name == "reports":
                frame.load_categories()
class DashboardFrame(ctk.CTkFrame):
    def __init__(self, parent, controller):
        super().__init__(parent)
        self.controller = controller
        title = ctk.CTkLabel(
            self,
            text="Панель управления",
            font=ctk.CTkFont(size=24, weight="bold")
        )
        title.pack(pady=20)
        all_inventory_btn = ctk.CTkButton(
            self,
            text="Весь инвентарь",
            command=self.show_all_inventory,
            width=300,
            height=40,
            font=ctk.CTkFont(size=16)
        )
        all_inventory_btn.pack(pady=10)
        self.stats_frame = ctk.CTkFrame(self)
        self.stats_frame.pack(fill="both", expand=True, padx=20, pady=20)
        self.inventory_frame = ctk.CTkFrame(self)
        self.update_dashboard()
    def update_dashboard(self):
        for widget in self.stats_frame.winfo_children():
            widget.destroy()
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM equipment")
        total = cursor.fetchone()[0]
        cursor.execute("SELECT COUNT(*) FROM equipment WHERE status='Активно'")
        active = cursor.fetchone()[0]
        cursor.execute("SELECT COUNT(*) FROM equipment WHERE status='Списано'")
        written_off = cursor.fetchone()[0]
        cursor.execute("SELECT COUNT(*) FROM categories")
        categories_count = cursor.fetchone()[0]
        conn.close()
        stats = [
            ("Всего оборудования", total),
            ("Активно", active),
            ("Списано", written_off),
            ("Категорий", categories_count)
        ]
        for i, (label, value) in enumerate(stats):
            card = ctk.CTkFrame(self.stats_frame, width=200, height=150)
            card.grid(row=i//2, column=i%2, padx=20, pady=20, sticky="nsew")
            ctk.CTkLabel(
                card,
                text=str(value),
                font=ctk.CTkFont(size=36, weight="bold")
            ).pack(pady=20)
            ctk.CTkLabel(
                card,
                text=label,
                font=ctk.CTkFont(size=16)
            ).pack()
        self.stats_frame.grid_columnconfigure(0, weight=1)
        self.stats_frame.grid_columnconfigure(1, weight=1)
    def show_all_inventory(self):
        self.stats_frame.pack_forget()
        self.inventory_frame.pack(fill="both", expand=True, padx=20, pady=20)
        if not hasattr(self, 'tree'):
            self.create_inventory_table()
        self.load_all_inventory()
    def create_inventory_table(self):
        back_btn = ctk.CTkButton(
            self.inventory_frame,
            text="← Назад к статистике",
            command=self.show_dashboard_stats,
            width=200,
            height=30
        )
        back_btn.pack(pady=5)
        columns = ("inv_number", "name", "model", "category", "date_in", "location", "responsible", "status", "description")
        self.tree = ttk.Treeview(self.inventory_frame, columns=columns, show="headings", height=15)
        headers = ["Инв. номер", "Наименование", "Модель", "Категория", "Дата прихода", "Местонахождение", "Ответственный", "Статус", "Описание"]
        for i, header in enumerate(headers):
            self.tree.heading(columns[i], text=header)
            self.tree.column(columns[i], width=120)
        self.tree.column("description", width=200)
        scrollbar = ttk.Scrollbar(self.inventory_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        self.tree.pack(side="left", fill="both", expand=True, padx=10, pady=10)
        scrollbar.pack(side="right", fill="y", pady=10)
    def load_all_inventory(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT e.inv_number, e.name, e.model, c.name, e.date_in, e.location, e.responsible, e.status, e.description
            FROM equipment e
            JOIN categories c ON e.category_id = c.id
            ORDER BY e.id
        """)
        rows = cursor.fetchall()
        conn.close()
        for row in rows:
            self.tree.insert("", "end", values=row)
    def show_dashboard_stats(self):
        self.inventory_frame.pack_forget()
        self.stats_frame.pack(fill="both", expand=True, padx=20, pady=20)
        self.update_dashboard()
class AddEquipmentFrame(ctk.CTkFrame):
    def __init__(self, parent, controller):
        super().__init__(parent)
        self.controller = controller
        title = ctk.CTkLabel(
            self,
            text="Добавить новое оборудование",
            font=ctk.CTkFont(size=24, weight="bold")
        )
        title.pack(pady=20)
        form_frame = ctk.CTkFrame(self)
        form_frame.pack(fill="both", expand=True, padx=40, pady=20)
        self.entries = {}
        fields = [
            ("inv_number", "Инвентарный номер *"),
            ("name", "Наименование *"),
            ("model", "Модель"),
            ("description", "Описание"),
            ("location", "Местонахождение *"),
            ("responsible", "ФИО ответственного *")
        ]
        for i, (key, label_text) in enumerate(fields):
            label = ctk.CTkLabel(form_frame, text=label_text, font=ctk.CTkFont(size=14))
            label.grid(row=i, column=0, padx=10, pady=10, sticky="w")
            if key == "description":
                entry = ctk.CTkTextbox(form_frame, height=60, width=400)
            else:
                entry = ctk.CTkEntry(form_frame, width=400)
            entry.grid(row=i, column=1, padx=10, pady=10)
            self.entries[key] = entry
        cat_label = ctk.CTkLabel(form_frame, text="Категория *", font=ctk.CTkFont(size=14))
        cat_label.grid(row=len(fields), column=0, padx=10, pady=10, sticky="w")
        self.category_var = ctk.StringVar()
        self.category_dropdown = ctk.CTkOptionMenu(form_frame, variable=self.category_var, width=400)
        self.category_dropdown.grid(row=len(fields), column=1, padx=10, pady=10)
        self.load_categories()
        save_btn = ctk.CTkButton(
            self,
            text="Сохранить",
            command=self.save_equipment,
            width=200,
            height=40,
            font=ctk.CTkFont(size=16)
        )
        save_btn.pack(pady=20)
    def load_categories(self):
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM categories ORDER BY name")
        categories = [row[0] for row in cursor.fetchall()]
        conn.close()
        if categories:
            self.category_var.set(categories[0])
            self.category_dropdown.configure(values=categories)
    def save_equipment(self):
        inv_number = self.entries["inv_number"].get().strip()
        name = self.entries["name"].get().strip()
        model = self.entries["model"].get().strip()
        description = self.entries["description"].get("1.0", "end-1c").strip()
        location = self.entries["location"].get().strip()
        responsible = self.entries["responsible"].get().strip()
        category = self.category_var.get()
        if not all([inv_number, name, location, responsible, category]):
            messagebox.showerror("Ошибка", "Все обязательные поля должны быть заполнены!")
            return
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT id FROM categories WHERE name=?", (category,))
        result = cursor.fetchone()
        if not result:
            messagebox.showerror("Ошибка", "Категория не найдена!")
            conn.close()
            return
        cat_id = result[0]
        date_in = datetime.now().strftime("%Y-%m-%d")
        try:
            cursor.execute("""
                INSERT INTO equipment (inv_number, name, model, category_id, location,
                                      responsible, date_in, description, status)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'Активно')
            """, (inv_number, name, model, cat_id, location, responsible, date_in, description))
            equip_id = cursor.lastrowid
            cursor.execute("""
                INSERT INTO history (equipment_id, action_type, date_action, new_location, comment)
                VALUES (?, 'Приход', ?, ?, 'Первичная регистрация')
            """, (equip_id, date_in, location))
            conn.commit()
            messagebox.showinfo("Успех", "Оборудование успешно добавлено!")
            for entry in self.entries.values():
                if isinstance(entry, ctk.CTkTextbox):
                    entry.delete("1.0", "end")
                else:
                    entry.delete(0, "end")
        except sqlite3.IntegrityError:
            messagebox.showerror("Ошибка", "Инвентарный номер уже существует!")
        finally:
            conn.close()
class MoveEquipmentFrame(ctk.CTkFrame):
    def __init__(self, parent, controller):
        super().__init__(parent)
        self.controller = controller
        title = ctk.CTkLabel(
            self,
            text="Перемещение оборудования",
            font=ctk.CTkFont(size=24, weight="bold")
        )
        title.pack(pady=20)
        form_frame = ctk.CTkFrame(self)
        form_frame.pack(fill="both", expand=True, padx=40, pady=20)
        search_frame = ctk.CTkFrame(form_frame)
        search_frame.grid(row=0, column=0, columnspan=2, pady=20, sticky="ew")
        ctk.CTkLabel(search_frame, text="Инвентарный номер:", font=ctk.CTkFont(size=14)).grid(row=0, column=0, padx=10, pady=10)
        self.inv_search = ctk.CTkEntry(search_frame, width=300)
        self.inv_search.grid(row=0, column=1, padx=10, pady=10)
        search_btn = ctk.CTkButton(search_frame, text="Найти", command=self.find_equipment, width=100)
        search_btn.grid(row=0, column=2, padx=10, pady=10)
        self.info_frame = ctk.CTkFrame(form_frame)
        self.info_frame.grid(row=1, column=0, columnspan=2, pady=20, sticky="nsew")
        self.info_labels = {}
        info_fields = ["Наименование", "Модель", "Категория", "Текущее местонахождение", "Ответственный"]
        for i, field in enumerate(info_fields):
            label = ctk.CTkLabel(self.info_frame, text=f"{field}:", font=ctk.CTkFont(size=14))
            label.grid(row=i, column=0, padx=10, pady=5, sticky="w")
            value_label = ctk.CTkLabel(self.info_frame, text="-", font=ctk.CTkFont(size=14))
            value_label.grid(row=i, column=1, padx=10, pady=5, sticky="w")
            self.info_labels[field] = value_label
        new_loc_label = ctk.CTkLabel(form_frame, text="Новое местонахождение:", font=ctk.CTkFont(size=14))
        new_loc_label.grid(row=2, column=0, padx=10, pady=10, sticky="w")
        self.new_location = ctk.CTkEntry(form_frame, width=400)
        self.new_location.grid(row=2, column=1, padx=10, pady=10)
        comment_label = ctk.CTkLabel(form_frame, text="Комментарий:", font=ctk.CTkFont(size=14))
        comment_label.grid(row=3, column=0, padx=10, pady=10, sticky="w")
        self.comment = ctk.CTkEntry(form_frame, width=400)
        self.comment.grid(row=3, column=1, padx=10, pady=10)
        self.move_btn = ctk.CTkButton(
            self,
            text="Переместить",
            command=self.move_equipment,
            width=200,
            height=40,
            font=ctk.CTkFont(size=16),
            state="disabled"
        )
        self.move_btn.pack(pady=20)
        self.current_equipment_id = None
    def find_equipment(self):
        inv_number = self.inv_search.get().strip()
        if not inv_number:
            messagebox.showerror("Ошибка", "Введите инвентарный номер!")
            return
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT e.id, e.name, e.model, c.name, e.location, e.responsible, e.status
            FROM equipment e
            JOIN categories c ON e.category_id = c.id
            WHERE e.inv_number = ?
        """, (inv_number,))
        result = cursor.fetchone()
        conn.close()
        if not result:
            messagebox.showerror("Ошибка", "Оборудование не найдено!")
            self.clear_info()
            return
        equip_id, name, model, category, location, responsible, status = result
        if status == "Списано":
            messagebox.showerror("Ошибка", "Оборудование списано и не может быть перемещено!")
            self.clear_info()
            return
        self.info_labels["Наименование"].configure(text=name)
        self.info_labels["Модель"].configure(text=model or "-")
        self.info_labels["Категория"].configure(text=category)
        self.info_labels["Текущее местонахождение"].configure(text=location)
        self.info_labels["Ответственный"].configure(text=responsible)
        self.current_equipment_id = equip_id
        self.move_btn.configure(state="normal")
    def clear_info(self):
        for label in self.info_labels.values():
            label.configure(text="-")
        self.current_equipment_id = None
        self.move_btn.configure(state="disabled")
    def move_equipment(self):
        if not self.current_equipment_id:
            return
        new_location = self.new_location.get().strip()
        comment = self.comment.get().strip()
        if not new_location:
            messagebox.showerror("Ошибка", "Укажите новое местонахождение!")
            return
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT location FROM equipment WHERE id=?", (self.current_equipment_id,))
        result = cursor.fetchone()
        if not result:
            conn.close()
            return
        old_location = result[0]
        date_now = datetime.now().strftime("%Y-%m-%d")
        cursor.execute("UPDATE equipment SET location=? WHERE id=?", (new_location, self.current_equipment_id))
        cursor.execute("""
            INSERT INTO history (equipment_id, action_type, date_action, old_location, new_location, comment)
            VALUES (?, 'Перемещение', ?, ?, ?, ?)
        """, (self.current_equipment_id, date_now, old_location, new_location, comment))
        conn.commit()
        conn.close()
        messagebox.showinfo("Успех", f"Оборудование перемещено из '{old_location}' в '{new_location}'")
        self.inv_search.delete(0, "end")
        self.new_location.delete(0, "end")
        self.comment.delete(0, "end")
        self.clear_info()
class WriteOffFrame(ctk.CTkFrame):
    def __init__(self, parent, controller):
        super().__init__(parent)
        self.controller = controller
        title = ctk.CTkLabel(
            self,
            text="Списание оборудования",
            font=ctk.CTkFont(size=24, weight="bold")
        )
        title.pack(pady=20)
        form_frame = ctk.CTkFrame(self)
        form_frame.pack(fill="both", expand=True, padx=40, pady=20)
        search_frame = ctk.CTkFrame(form_frame)
        search_frame.grid(row=0, column=0, columnspan=2, pady=20, sticky="ew")
        ctk.CTkLabel(search_frame, text="Инвентарный номер:", font=ctk.CTkFont(size=14)).grid(row=0, column=0, padx=10, pady=10)
        self.inv_search = ctk.CTkEntry(search_frame, width=300)
        self.inv_search.grid(row=0, column=1, padx=10, pady=10)
        search_btn = ctk.CTkButton(search_frame, text="Найти", command=self.find_equipment, width=100)
        search_btn.grid(row=0, column=2, padx=10, pady=10)
        self.info_frame = ctk.CTkFrame(form_frame)
        self.info_frame.grid(row=1, column=0, columnspan=2, pady=20, sticky="nsew")
        self.info_labels = {}
        info_fields = ["Наименование", "Модель", "Категория", "Местонахождение", "Ответственный", "Дата прихода"]
        for i, field in enumerate(info_fields):
            label = ctk.CTkLabel(self.info_frame, text=f"{field}:", font=ctk.CTkFont(size=14))
            label.grid(row=i, column=0, padx=10, pady=5, sticky="w")
            value_label = ctk.CTkLabel(self.info_frame, text="-", font=ctk.CTkFont(size=14))
            value_label.grid(row=i, column=1, padx=10, pady=5, sticky="w")
            self.info_labels[field] = value_label
        reason_label = ctk.CTkLabel(form_frame, text="Причина списания:", font=ctk.CTkFont(size=14))
        reason_label.grid(row=2, column=0, padx=10, pady=10, sticky="w")
        self.reason = ctk.CTkTextbox(form_frame, height=80, width=400)
        self.reason.grid(row=2, column=1, padx=10, pady=10)
        self.writeoff_btn = ctk.CTkButton(
            self,
            text="Списать",
            command=self.write_off_equipment,
            width=200,
            height=40,
            font=ctk.CTkFont(size=16),
            state="disabled"
        )
        self.writeoff_btn.pack(pady=20)
        self.current_equipment_id = None
    def find_equipment(self):
        inv_number = self.inv_search.get().strip()
        if not inv_number:
            messagebox.showerror("Ошибка", "Введите инвентарный номер!")
            return
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT e.id, e.name, e.model, c.name, e.location, e.responsible, e.date_in, e.status
            FROM equipment e
            JOIN categories c ON e.category_id = c.id
            WHERE e.inv_number = ?
        """, (inv_number,))
        result = cursor.fetchone()
        conn.close()
        if not result:
            messagebox.showerror("Ошибка", "Оборудование не найдено!")
            self.clear_info()
            return
        equip_id, name, model, category, location, responsible, date_in, status = result
        if status == "Списано":
            messagebox.showerror("Ошибка", "Оборудование уже списано!")
            self.clear_info()
            return
        self.info_labels["Наименование"].configure(text=name)
        self.info_labels["Модель"].configure(text=model or "-")
        self.info_labels["Категория"].configure(text=category)
        self.info_labels["Местонахождение"].configure(text=location)
        self.info_labels["Ответственный"].configure(text=responsible)
        self.info_labels["Дата прихода"].configure(text=date_in)
        self.current_equipment_id = equip_id
        self.writeoff_btn.configure(state="normal")
    def clear_info(self):
        for label in self.info_labels.values():
            label.configure(text="-")
        self.current_equipment_id = None
        self.writeoff_btn.configure(state="disabled")
    def write_off_equipment(self):
        if not self.current_equipment_id:
            return
        reason_text = self.reason.get("1.0", "end-1c").strip()
        if not reason_text:
            messagebox.showerror("Ошибка", "Укажите причину списания!")
            return
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT location FROM equipment WHERE id=?", (self.current_equipment_id,))
        result = cursor.fetchone()
        if not result:
            conn.close()
            return
        old_location = result[0]
        date_now = datetime.now().strftime("%Y-%m-%d")
        cursor.execute("UPDATE equipment SET status='Списано' WHERE id=?", (self.current_equipment_id,))
        cursor.execute("""
            INSERT INTO history (equipment_id, action_type, date_action, old_location, new_location, comment)
            VALUES (?, 'Списание', ?, ?, 'Архив', ?)
        """, (self.current_equipment_id, date_now, old_location, reason_text))
        conn.commit()
        conn.close()
        messagebox.showinfo("Успех", "Оборудование успешно списано!")
        self.inv_search.delete(0, "end")
        self.reason.delete("1.0", "end")
        self.clear_info()
class HistoryFrame(ctk.CTkFrame):
    def __init__(self, parent, controller):
        super().__init__(parent)
        self.controller = controller
        title = ctk.CTkLabel(
            self,
            text="История движений оборудования",
            font=ctk.CTkFont(size=24, weight="bold")
        )
        title.pack(pady=20)
        filter_frame = ctk.CTkFrame(self)
        filter_frame.pack(fill="x", padx=20, pady=10)
        ctk.CTkLabel(filter_frame, text="Инвентарный номер (оставьте пустым для всей истории):").pack(side="left", padx=10)
        self.inv_filter = ctk.CTkEntry(filter_frame, width=200)
        self.inv_filter.pack(side="left", padx=10)
        search_btn = ctk.CTkButton(filter_frame, text="Показать", command=self.load_history, width=100)
        search_btn.pack(side="left", padx=10)
        columns = ("date", "inv_number", "action", "old_location", "new_location", "comment")
        self.tree = ttk.Treeview(self, columns=columns, show="headings", height=20)
        self.tree.heading("date", text="Дата")
        self.tree.heading("inv_number", text="Инв. номер")
        self.tree.heading("action", text="Действие")
        self.tree.heading("old_location", text="Старая локация")
        self.tree.heading("new_location", text="Новая локация")
        self.tree.heading("comment", text="Комментарий")
        self.tree.column("date", width=100)
        self.tree.column("inv_number", width=120)
        self.tree.column("action", width=120)
        self.tree.column("old_location", width=150)
        self.tree.column("new_location", width=150)
        self.tree.column("comment", width=200)
        scrollbar = ttk.Scrollbar(self, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        self.tree.pack(side="left", fill="both", expand=True, padx=20, pady=20)
        scrollbar.pack(side="right", fill="y", pady=20)
        self.load_history()
    def load_history(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        inv_number = self.inv_filter.get().strip()
        conn = get_connection()
        cursor = conn.cursor()
        if inv_number:
            cursor.execute("""
                SELECT h.date_action, e.inv_number, h.action_type, h.old_location, h.new_location, h.comment
                FROM history h
                JOIN equipment e ON h.equipment_id = e.id
                WHERE e.inv_number = ?
                ORDER BY h.date_action DESC, h.id DESC
            """, (inv_number,))
        else:
            cursor.execute("""
                SELECT h.date_action, e.inv_number, h.action_type, h.old_location, h.new_location, h.comment
                FROM history h
                JOIN equipment e ON h.equipment_id = e.id
                ORDER BY h.date_action DESC, h.id DESC
                LIMIT 100
            """)
        rows = cursor.fetchall()
        conn.close()
        for row in rows:
            self.tree.insert("", "end", values=row)
class ReportsFrame(ctk.CTkFrame):
    def __init__(self, parent, controller):
        super().__init__(parent)
        self.controller = controller
        title = ctk.CTkLabel(
            self,
            text="Отчеты",
            font=ctk.CTkFont(size=24, weight="bold")
        )
        title.pack(pady=20)
        filter_frame = ctk.CTkFrame(self)
        filter_frame.pack(fill="x", padx=20, pady=10)
        ctk.CTkLabel(filter_frame, text="Тип отчета:").grid(row=0, column=0, padx=10, pady=10, sticky="w")
        self.report_type = ctk.StringVar(value="location")
        report_types = [
            ("По местонахождению", "location"),
            ("По дате прихода", "date"),
            ("По наименованию", "name"),
            ("По категории", "category")
        ]
        for i, (text, value) in enumerate(report_types):
            ctk.CTkRadioButton(filter_frame, text=text, variable=self.report_type,
                             value=value, command=self.on_report_type_change).grid(row=0, column=i+1, padx=10, pady=10)
        self.filter_frame = ctk.CTkFrame(self)
        self.filter_frame.pack(fill="x", padx=20, pady=10)
        self.setup_filters()
        generate_btn = ctk.CTkButton(self, text="Сгенерировать отчет", command=self.generate_report, width=200, height=40)
        generate_btn.pack(pady=10)
        columns = ("inv_number", "name", "model", "category", "date_in", "location", "responsible", "status", "description")
        self.tree = ttk.Treeview(self, columns=columns, show="headings", height=15)
        headers = ["Инв. номер", "Наименование", "Модель", "Категория", "Дата прихода", "Местонахождение", "Ответственный", "Статус", "Описание"]
        for i, header in enumerate(headers):
            self.tree.heading(columns[i], text=header)
            self.tree.column(columns[i], width=110)
        self.tree.column("description", width=200)
        self.tree.pack(fill="both", expand=True, padx=20, pady=20)
    def setup_filters(self):
        for widget in self.filter_frame.winfo_children():
            widget.destroy()
        report_type = self.report_type.get()
        if report_type == "location":
            ctk.CTkLabel(self.filter_frame, text="Местонахождение:").pack(side="left", padx=10)
            self.filter_entry = ctk.CTkEntry(self.filter_frame, width=300)
            self.filter_entry.pack(side="left", padx=10)
            self.filter_entry.insert(0, "")
        elif report_type == "date":
            ctk.CTkLabel(self.filter_frame, text="Дата с:").pack(side="left", padx=10)
            self.date_from = ctk.CTkEntry(self.filter_frame, width=150)
            self.date_from.pack(side="left", padx=10)
            self.date_from.insert(0, "2024-01-01")
            ctk.CTkLabel(self.filter_frame, text="по:").pack(side="left", padx=10)
            self.date_to = ctk.CTkEntry(self.filter_frame, width=150)
            self.date_to.pack(side="left", padx=10)
            self.date_to.insert(0, datetime.now().strftime("%Y-%m-%d"))
        elif report_type == "name":
            ctk.CTkLabel(self.filter_frame, text="Поиск по названию:").pack(side="left", padx=10)
            self.filter_entry = ctk.CTkEntry(self.filter_frame, width=300)
            self.filter_entry.pack(side="left", padx=10)
        elif report_type == "category":
            ctk.CTkLabel(self.filter_frame, text="Категория:").pack(side="left", padx=10)
            self.category_var = ctk.StringVar()
            self.filter_dropdown = ctk.CTkOptionMenu(self.filter_frame, variable=self.category_var, width=300)
            self.filter_dropdown.pack(side="left", padx=10)
            self.load_categories()
    def load_categories(self):
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM categories ORDER BY name")
        categories = ["Все"] + [row[0] for row in cursor.fetchall()]
        conn.close()
        if hasattr(self, 'category_var') and hasattr(self, 'filter_dropdown'):
            self.category_var.set(categories[0])
            self.filter_dropdown.configure(values=categories)
    def on_report_type_change(self):
        self.setup_filters()
    def generate_report(self):
        for item in self.tree.get_children():
            self.tree.delete(item)
        report_type = self.report_type.get()
        conn = get_connection()
        cursor = conn.cursor()
        query = """
            SELECT e.inv_number, e.name, e.model, c.name, e.date_in, e.location, e.responsible, e.status, e.description
            FROM equipment e
            JOIN categories c ON e.category_id = c.id
            WHERE 1=1
        """
        params = []
        if report_type == "location":
            location = self.filter_entry.get().strip()
            if location:
                query += " AND e.location LIKE ?"
                params.append(f"%{location}%")
        elif report_type == "date":
            date_from = self.date_from.get().strip()
            date_to = self.date_to.get().strip()
            if date_from:
                query += " AND e.date_in >= ?"
                params.append(date_from)
            if date_to:
                query += " AND e.date_in <= ?"
                params.append(date_to)
        elif report_type == "name":
            name = self.filter_entry.get().strip()
            if name:
                query += " AND e.name LIKE ?"
                params.append(f"%{name}%")
        elif report_type == "category":
            category = self.category_var.get()
            if category != "Все":
                query += " AND c.name = ?"
                params.append(category)
        query += " ORDER BY e.id"
        cursor.execute(query, params)
        rows = cursor.fetchall()
        conn.close()
        for row in rows:
            self.tree.insert("", "end", values=row)
        if not rows:
            messagebox.showinfo("Информация", "По заданным критериям ничего не найдено")
class ImportExportFrame(ctk.CTkFrame):
    def __init__(self, parent, controller):
        super().__init__(parent)
        self.controller = controller
        title = ctk.CTkLabel(
            self,
            text="Импорт и экспорт данных",
            font=ctk.CTkFont(size=24, weight="bold")
        )
        title.pack(pady=20)
        export_frame = ctk.CTkFrame(self)
        export_frame.pack(fill="x", padx=40, pady=20)
        ctk.CTkLabel(export_frame, text="Экспорт в Excel", font=ctk.CTkFont(size=18, weight="bold")).pack(pady=10)
        export_btn = ctk.CTkButton(
            export_frame,
            text="Экспортировать все данные",
            command=self.export_to_excel,
            width=300,
            height=40
        )
        export_btn.pack(pady=10)
        import_frame = ctk.CTkFrame(self)
        import_frame.pack(fill="x", padx=40, pady=20)
        ctk.CTkLabel(import_frame, text="Импорт из Excel", font=ctk.CTkFont(size=18, weight="bold")).pack(pady=10)
        import_btn = ctk.CTkButton(
            import_frame,
            text="Импортировать из файла",
            command=self.import_from_excel,
            width=300,
            height=40
        )
        import_btn.pack(pady=10)
        info_text = """
        Формат файла для импорта (Excel):
        Столбцы: Инв.номер | Наименование | Модель | Категория | Местонахождение | Ответственный | Дата прихода | Статус | Описание

        Пример:
        INV-001 | Моноблок | Lenovo ThinkCentre | ПК | Ауд. 313 | Иванов И.И. | 2024-01-15 | Активно | Новый
        """
        info_label = ctk.CTkLabel(
            self,
            text=info_text,
            font=ctk.CTkFont(size=12),
            justify="left"
        )
        info_label.pack(pady=20, padx=40)
    def export_to_excel(self):
        filename = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile="inventory_export.xlsx"
        )
        if not filename:
            return
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT e.inv_number, e.name, e.model, c.name as category, e.location,
                   e.responsible, e.date_in, e.status, e.description
            FROM equipment e
            JOIN categories c ON e.category_id = c.id
            ORDER BY e.id
        """)
        rows = cursor.fetchall()
        conn.close()
        if not rows:
            messagebox.showerror("Ошибка", "Нет данных для экспорта!")
            return
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Инвентаризация"
        headers = ["Инв. номер", "Наименование", "Модель", "Категория", "Местонахождение",
                   "Ответственный", "Дата прихода", "Статус", "Описание"]
        ws.append(headers)
        for row in rows:
            ws.append(row)
        wb.save(filename)
        messagebox.showinfo("Успех", f"Данные экспортированы в файл: {filename}")
    def import_from_excel(self):
        filename = filedialog.askopenfilename(
            filetypes=[("Excel files", "*.xlsx")],
            title="Выберите файл для импорта"
        )
        if not filename:
            return
        try:
            wb = openpyxl.load_workbook(filename)
            ws = wb.active
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            if not rows or rows[0][0] is None:
                messagebox.showerror("Ошибка", "File пуст или не содержит данных!")
                return
            conn = get_connection()
            cursor = conn.cursor()
            cursor.execute("SELECT id, name FROM categories")
            cat_map = {name.lower(): id for id, name in cursor.fetchall()}
            imported_count = 0
            skipped_count = 0
            for row in rows:
                if len(row) < 5:
                    skipped_count += 1
                    continue
                inv_num = str(row[0]).strip() if row[0] else ""
                name = str(row[1]).strip() if row[1] else ""
                model = str(row[2]).strip() if row[2] else ""
                cat_name = str(row[3]).strip().lower() if row[3] else "прочее"
                location = str(row[4]).strip() if row[4] else "Не указано"
                responsible = str(row[5]).strip() if row[5] else "Не указано"
                date_in = str(row[6]).strip() if row[6] else datetime.now().strftime("%Y-%m-%d")
                status = str(row[7]).strip() if row[7] else "Активно"
                desc = str(row[8]).strip() if len(row) > 8 and row[8] else ""
                if not inv_num or not name:
                    skipped_count += 1
                    continue
                cat_id = cat_map.get(cat_name, None)
                if not cat_id:
                    cursor.execute("INSERT INTO categories (name) VALUES (?)", (cat_name.capitalize(),))
                    cat_id = cursor.lastrowid
                    cat_map[cat_name] = cat_id
                cursor.execute("SELECT id FROM equipment WHERE inv_number=?", (inv_num,))
                if cursor.fetchone():
                    skipped_count += 1
                    continue
                cursor.execute("""
                    INSERT INTO equipment (inv_number, name, model, category_id, location, responsible, date_in, status, description)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (inv_num, name, model, cat_id, location, responsible, date_in, status, desc))
                equip_id = cursor.lastrowid
                cursor.execute("""
                    INSERT INTO history (equipment_id, action_type, date_action, new_location, comment)
                    VALUES (?, 'Импорт из Excel', ?, ?, 'Массовая загрузка')
                """, (equip_id, date_in, location))
                imported_count += 1
            conn.commit()
            conn.close()
            messagebox.showinfo("Успех", f"Импорт завершен!\nДобавлено: {imported_count}\nПропущено: {skipped_count}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Ошибка при импорте: {str(e)}")
def main():
    app = InventoryApp()
    app.mainloop()
if __name__ == "__main__":
    main()